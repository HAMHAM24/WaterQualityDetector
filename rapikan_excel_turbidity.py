"""Rapikan workbook data turbidity aktif dan tandai workbook lama sebagai arsip.

Jalankan dari root proyek:
    python rapikan_excel_turbidity.py

Memerlukan openpyxl:
    python -m pip install openpyxl
"""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Series, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
ACTIVE_BOOK = ROOT / "turbidity compare.xlsx"
ARCHIVE_BOOK = ROOT / "Perbandingan_Turbidity.xlsx"

DATA = [
    (1, "Full aquades", "Penuh", 0.00, 710, 0.00),
    (2, "Aquades + Formazin", 14.70, 0.30, 718, 9.44),
    (3, "Aquades + Formazin", 14.40, 0.60, 726, 19.47),
    (4, "Aquades + Formazin", 14.01, 0.90, 732, 28.97),
    (5, "Aquades + Formazin", 13.80, 1.20, 742, 38.72),
    (6, "Aquades + Formazin", 13.31, 1.69, 748, 47.45),
    (7, "Aquades + Formazin", 9.00, 6.00, 852, 177.30),
    (8, "Full Formazin", 0.00, "Penuh", 1084, 468.00),
]

OLD_READINGS = [
    (1, 710, 211.5, 0.55, "Tidak digunakan untuk regresi"),
    (2, 718, 206.9, None, "Tidak digunakan untuk regresi"),
    (3, 726, 202.0, None, "Tidak digunakan untuk regresi"),
    (4, 732, 197.4, None, "Tidak digunakan untuk regresi"),
    (5, 742, 192.6, None, "Tidak digunakan untuk regresi"),
    (6, 748, 188.4, None, "Tidak digunakan untuk regresi"),
    (7, 852, 192.0, None, "Tidak digunakan untuk regresi"),
    (9, 1084, 0.0, 0.88, "Tidak digunakan untuk regresi"),
    (10, 637, 257.0, 0.51, "Sensor tidak dicelup; data diagnostik"),
]

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="B7C9D6")


def style_title(ws, title, end_column):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    cell = ws.cell(1, 1, title)
    cell.font = Font(bold=True, color=WHITE, size=14)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center")


def style_header(row):
    for cell in row:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)


def style_table(ws, start_row, end_row, end_column):
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=1, max_col=end_column):
        for cell in row:
            cell.border = Border(bottom=THIN_GRAY)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def set_widths(ws, widths):
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width


def create_active_workbook():
    workbook = load_workbook(ACTIVE_BOOK)
    original = workbook["Data Asli"] if "Data Asli" in workbook.sheetnames else workbook[workbook.sheetnames[0]]

    # Script dapat dijalankan ulang tanpa menggandakan sheet sumber atau sheet
    # hasil olahan. Sisakan satu sheet mentah sebagai audit trail.
    for sheet_name in list(workbook.sheetnames):
        if sheet_name.startswith("Data Asli") and workbook[sheet_name] is not original:
            del workbook[sheet_name]
    if original.title != "Data Asli":
        original.title = "Data Asli"
    original.sheet_view.showGridLines = False

    if "README" in workbook.sheetnames:
        del workbook["README"]

    for sheet_name in ("Data Kalibrasi", "Regresi", "Pembacaan Lama"):
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]

    readme_ws = workbook.create_sheet("README", 0)
    readme_ws.sheet_view.showGridLines = False
    style_title(readme_ws, "Workbook Kalibrasi Turbidity Aktif", 6)
    readme_ws.merge_cells("A3:F3")
    readme_ws["A3"] = "Gunakan workbook ini untuk data Lab Bante terbaru. Jangan gunakan workbook Perbandingan_Turbidity.xlsx untuk menghitung koefisien."
    readme_ws["A3"].alignment = Alignment(wrap_text=True)
    readme_ws["A3"].fill = PatternFill("solid", fgColor=YELLOW)
    readme_ws.row_dimensions[3].height = 32
    readme_ws["A5"] = "Urutan penggunaan"
    readme_ws["A5"].font = Font(bold=True, color=WHITE)
    readme_ws["A5"].fill = PatternFill("solid", fgColor=BLUE)
    steps = [
        "Catat ADC stabil pada Live Turbidity dan NTU dari Lab Bante untuk sampel yang sama.",
        "Tambahkan pasangan data baru pada Data Kalibrasi dan kalibrasi_turbidity.py.",
        "Jalankan: python kalibrasi_turbidity.py --no-plot.",
        "Salin slope dan intercept hasil script ke main/config.h, lalu build dan upload firmware.",
    ]
    for row, step in enumerate(steps, start=6):
        readme_ws.cell(row, 1, f"{row - 5}.")
        readme_ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        readme_ws.cell(row, 2, step)
        readme_ws.cell(row, 2).alignment = Alignment(wrap_text=True)
        readme_ws.row_dimensions[row].height = 30
    readme_ws["A12"] = "Status pembacaan"
    readme_ws["A12"].font = Font(bold=True, color=WHITE)
    readme_ws["A12"].fill = PatternFill("solid", fgColor=BLUE)
    statuses = [
        ("BAWAH RENTANG", "ADC < 710; hasil dijepit ke 0 NTU."),
        ("TERKALIBRASI", "ADC 710-1084; rentang 0-468 NTU sudah diuji."),
        ("ESTIMASI >468", "ADC > 1084; hasil belum diverifikasi Lab Bante."),
        ("ERROR", "ADC mentah 0 atau 4095; periksa sensor/rangkaian."),
    ]
    for row, (status, description) in enumerate(statuses, start=13):
        readme_ws.cell(row, 1, status)
        readme_ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        readme_ws.cell(row, 2, description)
        readme_ws.cell(row, 2).alignment = Alignment(wrap_text=True)
    style_table(readme_ws, 6, 16, 6)
    set_widths(readme_ws, [20, 26, 18, 18, 18, 18])

    data_ws = workbook.create_sheet("Data Kalibrasi", 1)
    data_ws.sheet_view.showGridLines = False
    style_title(data_ws, "Data Kalibrasi Turbidity - Referensi Lab Bante", 8)
    data_ws.merge_cells("A2:H2")
    data_ws["A2"] = (
        "Gunakan hanya ADC Alatku dan NTU Lab Bante untuk regresi. "
        "Rentang tervalidasi: 0-468 NTU."
    )
    data_ws["A2"].alignment = Alignment(wrap_text=True)
    data_ws["A2"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    data_ws["A2"].font = Font(italic=True)
    headers = [
        "No.", "Sampel", "Aquades (mL)", "Formazin 500 NTU (mL)",
        "ADC Alatku", "NTU Lab Bante", "Prediksi Regresi (NTU)", "Residual (Lab - Prediksi)",
    ]
    for index, header in enumerate(headers, start=1):
        data_ws.cell(4, index, header)
    style_header(data_ws[4])
    for row_index, data in enumerate(DATA, start=5):
        for column, value in enumerate(data, start=1):
            data_ws.cell(row_index, column, value)
        data_ws.cell(row_index, 7, f"=$K$5*E{row_index}+$K$6")
        data_ws.cell(row_index, 8, f"=F{row_index}-G{row_index}")
    style_table(data_ws, 5, 12, 8)
    for row in range(5, 13):
        data_ws.cell(row, 3).number_format = "0.00"
        data_ws.cell(row, 4).number_format = "0.00"
        for column in (6, 7, 8):
            data_ws.cell(row, column).number_format = "0.00"

    data_ws["J4"] = "Koefisien Firmware"
    data_ws["J4"].font = Font(bold=True, color=WHITE)
    data_ws["J4"].fill = PatternFill("solid", fgColor=BLUE)
    data_ws["K4"] = "Nilai"
    data_ws["K4"].font = Font(bold=True, color=WHITE)
    data_ws["K4"].fill = PatternFill("solid", fgColor=BLUE)
    data_ws["J5"] = "Slope"
    data_ws["K5"] = 1.2516445883017
    data_ws["J6"] = "Intercept"
    data_ws["K6"] = -888.878830170042
    data_ws["J7"] = "R kuadrat"
    data_ws["K7"] = 0.999975641836792
    data_ws["J8"] = "RMSE (NTU)"
    data_ws["K8"] = 0.735
    data_ws["J9"] = "ADC minimum"
    data_ws["K9"] = 710
    data_ws["J10"] = "ADC maksimum"
    data_ws["K10"] = 1084
    data_ws["J11"] = "NTU minimum"
    data_ws["K11"] = 0
    data_ws["J12"] = "NTU maksimum"
    data_ws["K12"] = 468
    style_table(data_ws, 4, 12, 11)
    for row in range(5, 9):
        data_ws.cell(row, 11).number_format = "0.000000"
    set_widths(data_ws, [7, 24, 15, 22, 14, 16, 22, 24, 3, 22, 16])
    data_ws.freeze_panes = "A5"
    data_ws.auto_filter.ref = "A4:H12"

    regression_ws = workbook.create_sheet("Regresi", 1)
    regression_ws.sheet_view.showGridLines = False
    style_title(regression_ws, "Ringkasan Regresi Linear Turbidity", 6)
    regression_ws.merge_cells("A3:F3")
    regression_ws["A3"] = "Model: NTU = slope x ADC + intercept"
    regression_ws["A3"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    regression_ws["A3"].font = Font(bold=True)
    regression_ws["A5"] = "Parameter"
    regression_ws["B5"] = "Nilai"
    style_header(regression_ws[5][0:2])
    metrics = [
        ("Jumlah titik (n)", 8),
        ("Jumlah ADC", 6312),
        ("Jumlah NTU", 789.35),
        ("Rata-rata ADC (x_bar)", 789),
        ("Rata-rata NTU (y_bar)", 98.66875),
        ("Sxy", 141916.47),
        ("Sxx", 113384),
        ("Slope (m)", 1.2516445883017),
        ("Intercept (b)", -888.878830170042),
        ("R kuadrat", 0.999975641836792),
        ("RMSE (NTU)", 0.735),
    ]
    for row_index, (label, value) in enumerate(metrics, start=6):
        regression_ws.cell(row_index, 1, label)
        regression_ws.cell(row_index, 2, value)
        regression_ws.cell(row_index, 2).number_format = "0.000000000000"
    style_table(regression_ws, 6, 16, 2)
    regression_ws["A18"] = "Rumus firmware"
    regression_ws["A18"].font = Font(bold=True, color=WHITE)
    regression_ws["A18"].fill = PatternFill("solid", fgColor=BLUE)
    regression_ws.merge_cells("B18:F18")
    regression_ws["B18"] = "NTU = 1.251645 x ADC - 888.878830"
    regression_ws["B18"].fill = PatternFill("solid", fgColor=GREEN)
    regression_ws["A20"] = "Batas klaim"
    regression_ws["A20"].font = Font(bold=True, color=WHITE)
    regression_ws["A20"].fill = PatternFill("solid", fgColor=BLUE)
    regression_ws.merge_cells("B20:F20")
    regression_ws["B20"] = "Terkalibrasi pada ADC 710-1084 atau 0-468 NTU. Di atasnya hanya estimasi."
    regression_ws["B20"].fill = PatternFill("solid", fgColor=YELLOW)
    regression_ws["B20"].alignment = Alignment(wrap_text=True)
    regression_ws.row_dimensions[20].height = 32
    chart = ScatterChart()
    chart.title = "Kalibrasi ADC terhadap NTU Lab Bante"
    chart.x_axis.title = "ADC Alatku"
    chart.y_axis.title = "NTU Lab Bante"
    x_values = Reference(data_ws, min_col=5, min_row=5, max_row=12)
    y_values = Reference(data_ws, min_col=6, min_row=5, max_row=12)
    series = Series(y_values, x_values, title="Data Lab Bante")
    series.marker.symbol = "circle"
    series.graphicalProperties.line.noFill = True
    chart.series.append(series)
    regression_ws.add_chart(chart, "D5")
    set_widths(regression_ws, [28, 22, 3, 18, 18, 18])

    old_ws = workbook.create_sheet("Pembacaan Lama")
    old_ws.sheet_view.showGridLines = False
    style_title(old_ws, "Pembacaan Alat Sebelum Regresi - Tidak Dipakai untuk Kalibrasi", 5)
    old_ws.merge_cells("A2:E2")
    old_ws["A2"] = "Nilai ini dipertahankan sebagai riwayat. Jangan gunakan sebagai NTU referensi atau untuk menghitung koefisien firmware."
    old_ws["A2"].fill = PatternFill("solid", fgColor=RED)
    old_ws["A2"].alignment = Alignment(wrap_text=True)
    headers = ["No. Asal", "ADC Alatku", "NTU Alatku Lama", "Volt Alatku", "Catatan"]
    for index, header in enumerate(headers, start=1):
        old_ws.cell(4, index, header)
    style_header(old_ws[4])
    for row_index, data in enumerate(OLD_READINGS, start=5):
        for column, value in enumerate(data, start=1):
            old_ws.cell(row_index, column, value)
    style_table(old_ws, 5, 13, 5)
    for row in range(5, 14):
        old_ws.cell(row, 3).number_format = "0.0"
        old_ws.cell(row, 4).number_format = "0.00"
    set_widths(old_ws, [12, 14, 18, 14, 42])
    old_ws.freeze_panes = "A5"

    workbook.save(ACTIVE_BOOK)


def mark_archive_workbook():
    workbook = load_workbook(ARCHIVE_BOOK)
    ws = workbook.active
    # Bersihkan banner dari eksekusi sebelumnya agar script idempotent.
    while ws["A1"].value == "ARSIP PENGUJIAN LAMA - BUKAN DASAR KALIBRASI FIRMWARE AKTIF":
        ws.delete_rows(1, 2)
    ws.insert_rows(1, 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws["A1"] = "ARSIP PENGUJIAN LAMA - BUKAN DASAR KALIBRASI FIRMWARE AKTIF"
    ws["A1"].font = Font(bold=True, color=WHITE, size=12)
    ws["A1"].fill = PatternFill("solid", fgColor="9C0006")
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    ws["A2"] = "Gunakan turbidity compare.xlsx dan sampel_baru_turbidity.md untuk regresi Lab Bante terbaru."
    ws["A2"].fill = PatternFill("solid", fgColor=RED)
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A5"
    workbook.save(ARCHIVE_BOOK)


def main():
    create_active_workbook()
    mark_archive_workbook()
    print("Workbook turbidity aktif dan arsip berhasil dirapikan.")


if __name__ == "__main__":
    main()
